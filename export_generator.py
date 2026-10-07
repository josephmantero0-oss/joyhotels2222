import os
import io
import json
import sqlite3
from datetime import datetime
import zipfile
from PIL import Image

def fix_path(path):
    if path and path.startswith('uploads/'):
        return os.path.join('static', path)
    return path

def get_thumbnail(path, max_width=100, max_height=100):
    path = fix_path(path)
    if not path or not os.path.exists(path):
        return None
    try:
        img = Image.open(path)
        # Convert RGBA to RGB for JPEG
        if img.mode in ('RGBA', 'P'): 
            img = img.convert('RGB')
        img.thumbnail((max_width, max_height))
        
        output = io.BytesIO()
        img.save(output, format="JPEG", quality=85)
        output.seek(0)
        return output
    except Exception as e:
        print(f"Error loading image {path}: {e}")
        return None

def generate_excel(bookings, hotel_name):
    import openpyxl
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.styles import Font, Alignment
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Guest History"
    
    headers = ['ID', 'Guest Name', 'Phone', 'Room', 'Check-in', 'Check-out', 'Total (Rs)', 'Status', 'ID Photo']
    ws.append(headers)
    
    for col in range(1, 10):
        ws.cell(row=1, column=col).font = Font(bold=True)
        ws.cell(row=1, column=col).alignment = Alignment(horizontal='center')
        
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['I'].width = 20
        
    for idx, b in enumerate(bookings, start=2):
        total = (b['amount'] or 0) + (b['extra_charges'] or 0) + (b['food'] or 0) + \
                (b['laundry'] or 0) + (b['water_bottle'] or 0) + (b['car_wash'] or 0)
        
        raw_members = b['members'] or '[]'
        member_str = ""
        try:
            members_list = json.loads(raw_members)
            if isinstance(members_list, list):
                member_str = ', '.join([m if isinstance(m, str) else (m.get('name') or '') for m in members_list if m])
        except:
            pass
            
        all_names = b['guest_name'] or ''
        if member_str:
            all_names += f"\n({member_str})"
            
        row_data = [
            b['id'], all_names, b['phone'], b['room_id'], 
            b['checkin_time'], b['checkout_time'], total, b['status']
        ]
        
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=idx, column=col_idx)
            cell.value = value
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            
        id_photos = (b['id_photo'] or '').split(',')
        if id_photos and id_photos[0]:
            photo_path = fix_path(id_photos[0])
            if os.path.exists(photo_path):
                img_io = get_thumbnail(photo_path, max_width=120, max_height=80)
                if img_io:
                    try:
                        img = OpenpyxlImage(img_io)
                        ws.row_dimensions[idx].height = 70
                        ws.add_image(img, f"I{idx}")
                    except Exception as e:
                        pass
                        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_word(bookings, hotel_name):
    from docx import Document
    from docx.shared import Inches
    
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        
    doc.add_heading(f'{hotel_name} - Guest History Backup', 0)
    doc.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    
    table = doc.add_table(rows=1, cols=8)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    headers = ['ID', 'Guest Name', 'Phone', 'Room', 'Check-in', 'Total', 'Status', 'ID Photo']
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        
    for b in bookings:
        row_cells = table.add_row().cells
        total = (b['amount'] or 0) + (b['extra_charges'] or 0) + (b['food'] or 0) + \
                (b['laundry'] or 0) + (b['water_bottle'] or 0) + (b['car_wash'] or 0)
                
        raw_members = b['members'] or '[]'
        member_str = ""
        try:
            members_list = json.loads(raw_members)
            if isinstance(members_list, list):
                member_str = ', '.join([m if isinstance(m, str) else (m.get('name') or '') for m in members_list if m])
        except:
            pass
            
        all_names = b['guest_name'] or ''
        if member_str:
            all_names += f"\n({member_str})"
            
        row_cells[0].text = str(b['id'])
        row_cells[1].text = all_names
        row_cells[2].text = str(b['phone'] or '')
        row_cells[3].text = str(b['room_id'] or '')
        row_cells[4].text = str(b['checkin_time'] or '')
        row_cells[5].text = str(total)
        row_cells[6].text = str(b['status'] or '')
        
        id_photos = (b['id_photo'] or '').split(',')
        if id_photos and id_photos[0]:
            photo_path = fix_path(id_photos[0])
            if os.path.exists(photo_path):
                img_io = get_thumbnail(photo_path, max_width=100, max_height=80)
                if img_io:
                    try:
                        paragraph = row_cells[7].paragraphs[0]
                        run = paragraph.add_run()
                        run.add_picture(img_io, width=Inches(1.0))
                    except:
                        pass
                        
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output

def generate_pdf(bookings, hotel_name):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Image, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter), rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    
    styles = getSampleStyleSheet()
    elements = []
    
    elements.append(Paragraph(f"<b>{hotel_name} - Guest History Backup</b>", styles['Heading1']))
    elements.append(Paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    data = [['ID', 'Guest Name', 'Phone', 'Room', 'Check-in', 'Total', 'Status', 'ID Photo']]
    
    for b in bookings:
        total = (b['amount'] or 0) + (b['extra_charges'] or 0) + (b['food'] or 0) + \
                (b['laundry'] or 0) + (b['water_bottle'] or 0) + (b['car_wash'] or 0)
                
        raw_members = b['members'] or '[]'
        member_str = ""
        try:
            members_list = json.loads(raw_members)
            if isinstance(members_list, list):
                member_str = ', '.join([m if isinstance(m, str) else (m.get('name') or '') for m in members_list if m])
        except:
            pass
            
        all_names = b['guest_name'] or ''
        if member_str:
            all_names += f"\n({member_str})"
            
        photo_element = ""
        id_photos = (b['id_photo'] or '').split(',')
        if id_photos and id_photos[0]:
            photo_path = fix_path(id_photos[0])
            if os.path.exists(photo_path):
                img_io = get_thumbnail(photo_path, max_width=80, max_height=60)
                if img_io:
                    try:
                        photo_element = Image(img_io, width=80, height=60)
                    except:
                        pass
                        
        row = [
            str(b['id']),
            Paragraph(all_names.replace('\n', '<br/>'), styles['Normal']),
            str(b['phone'] or ''),
            str(b['room_id'] or ''),
            str(b['checkin_time'] or ''),
            str(total),
            str(b['status'] or ''),
            photo_element
        ]
        data.append(row)
        
    colWidths = [30, 150, 80, 50, 100, 60, 60, 90]
    t = Table(data, colWidths=colWidths)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4f46e5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    
    elements.append(t)
    doc.build(elements)
    
    output.seek(0)
    return output

def generate_photo_word(bookings, hotel_name):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
    doc.add_heading(f'{hotel_name} - Guest ID Photos', 0)
    doc.add_paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Sr No'
    hdr_cells[1].text = 'Guest Name'
    hdr_cells[2].text = 'ID Photo'
    
    sr_no = 1
    for b in bookings:
        raw_members = b['members'] or '[]'
        member_str = ""
        try:
            members_list = json.loads(raw_members)
            if isinstance(members_list, list):
                member_str = ', '.join([m if isinstance(m, str) else (m.get('name') or '') for m in members_list if m])
        except:
            pass
            
        all_names = b['guest_name'] or ''
        if member_str:
            all_names += f"\n({member_str})"
        
        # Get all photos for this guest
        id_photos = [p.strip() for p in (b['id_photo'] or '').split(',') if p.strip()]
        
        row_cells = table.add_row().cells
        row_cells[0].text = str(sr_no)
        row_cells[1].text = all_names
        
        # Add all photos into the photo cell
        for photo_rel in id_photos:
            photo_path = fix_path(photo_rel)
            if os.path.exists(photo_path):
                img_io = get_thumbnail(photo_path, max_width=400, max_height=350)
                if img_io:
                    try:
                        paragraph = row_cells[2].paragraphs[0] if row_cells[2].paragraphs else row_cells[2].add_paragraph()
                        run = paragraph.add_run()
                        run.add_picture(img_io, width=Inches(3.0))
                        # Add a line break after photo if there are multiple
                        run.add_break()
                    except:
                        pass
        
        sr_no += 1
                        
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output

def generate_multi_backup(hotel_id, hotel_name):
    conn = sqlite3.connect('hotel.db')
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute('''SELECT * FROM bookings WHERE hotel_id = ? ORDER BY id DESC''', (hotel_id,))
    bookings = [dict(row) for row in c.fetchall()]
    conn.close()
    
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        excel_io = generate_excel(bookings, hotel_name)
        zf.writestr(f"Guest_History_{hotel_name}.xlsx", excel_io.getvalue())
        
        word_io = generate_word(bookings, hotel_name)
        zf.writestr(f"Guest_History_{hotel_name}.docx", word_io.getvalue())
        
        pdf_io = generate_pdf(bookings, hotel_name)
        zf.writestr(f"Guest_History_{hotel_name}.pdf", pdf_io.getvalue())
        
        photos_word_io = generate_photo_word(bookings, hotel_name)
        zf.writestr(f"Photos_{hotel_name}.docx", photos_word_io.getvalue())
        
    zip_buffer.seek(0)
    return zip_buffer
