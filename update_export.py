import os
import re

with open('export_generator.py', 'r', encoding='utf-8') as f:
    content = f.read()

new_func = '''def generate_photo_word(bookings, hotel_name):
    from docx import Document
    from docx.shared import Inches
    import json
    import os
    import io
    
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
    doc.add_heading(f'{hotel_name} - Guest Photos', 0)
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Sr No'
    hdr_cells[1].text = 'Guest Name'
    hdr_cells[2].text = 'ID Photo'
    
    sr_no = 1
    for b in bookings:
        row_cells = table.add_row().cells
        
        raw_members = b['members'] or '[]'
        member_str = ""
        try:
            members_list = json.loads(raw_members)
            if isinstance(members_list, list):
                member_str = ', '.join([m if isinstance(m, str) else (m.get('name') or '') for m in members_list if m])
        except:
            pass
            
        all_names = b['guest_name'] or ''
        if member_str:
            all_names += f"\\n({member_str})"
            
        row_cells[0].text = str(sr_no)
        row_cells[1].text = all_names
        
        id_photos = (b['id_photo'] or '').split(',')
        if id_photos and id_photos[0]:
            photo_path = fix_path(id_photos[0])
            if os.path.exists(photo_path):
                # Larger size for this Word document
                img_io = get_thumbnail(photo_path, max_width=300, max_height=250)
                if img_io:
                    try:
                        paragraph = row_cells[2].paragraphs[0]
                        run = paragraph.add_run()
                        run.add_picture(img_io, width=Inches(3.0))
                    except:
                        pass
        sr_no += 1
                        
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output

def generate_multi_backup(hotel_id, hotel_name):
    import sqlite3
    import zipfile
    import io
    
    conn = sqlite3.connect('hotel.db')
    conn.row_factory = sqlite3.Row
    c = conn.cursor()
    c.execute(\'\'\'SELECT * FROM bookings WHERE hotel_id = ? ORDER BY id DESC\'\'\', (hotel_id,))
    bookings = [dict(row) for row in c.fetchall()]
    conn.close()
    
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        excel_io = generate_excel(bookings, hotel_name)
        zf.writestr(f"Guest_History_{hotel_name}.xlsx", excel_io.getvalue())
        
        word_io = generate_word(bookings, hotel_name)
        zf.writestr(f"Guest_History_{hotel_name}.docx", word_io.getvalue())
        
        pdf_io = generate_pdf(bookings, hotel_name)
        zf.writestr(f"Guest_History_{hotel_name}.pdf", pdf_io.getvalue())
        
        photos_word_io = generate_photo_word(bookings, hotel_name)
        zf.writestr(f"Photos_{hotel_name}.docx", photos_word_io.getvalue())
        
    zip_buffer.seek(0)
    return zip_buffer
'''

# Replace from `def generate_photo_pdf` to the end of the file
new_content = re.sub(r'def generate_photo_pdf\(.*', new_func, content, flags=re.DOTALL)

with open('export_generator.py', 'w', encoding='utf-8') as f:
    f.write(new_content)
